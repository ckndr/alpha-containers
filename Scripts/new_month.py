"""
new_month.py
Tubex — Monthly Workbook Rollover Automation
──────────────────────────────────────────────────────────
Rolls the active Tubex workbook over to a new month.

Usage:
  python new_month.py --month Oct --year 2026 [--dry-run]
  python new_month.py -m Oct -y 2026

Steps performed:
1. Confirm no Tubex_<NewMon><NewYY>.xlsx already exists in root (aborts if found).
2. Copy current active workbook (via get_active_tubex_file) to new filename in root.
3. In new copy's MRP sheet: for every currently active TUBE and PET order row,
   compute remaining = Required Qty - Produced.
   - If remaining > 0: set Required Qty = remaining, Produced = 0, keep JOF and PID.
   - If remaining <= 0: remove row entirely (finished order).
   Row removal and formula reflow reuse add_order.py logic (rearrange_mrp_orders).
4. Clear Production_Log to just its header rows (rows 1-2).
   Clear FG Stock to just its header rows (rows 1-3).
5. Preserve BOM, Product_Catalog, Inventory formulas/structure, and MRP material rows.
6. Reset Dashboard orders/dispatch and run sort_dashboard.py against the new workbook
   so Dashboard rebuilds cleanly from fresh MRP state (Produced=0, Dispatch=0).
7. Move OLD month workbook into "Tubex Records\\".
8. Print summary of carried-forward and finished orders, and file locations.
9. --dry-run operates entirely inside a temporary sandbox without modifying live files.

Author: Sikander / Antigravity
"""

import os
import sys
import glob
import re
import shutil
import argparse
import subprocess
import tempfile
from datetime import datetime
import openpyxl

# Add Scripts folder to sys.path so helper modules can be imported
SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)
# Fallback to D:\Alpha\Scripts if needed
fallback_scripts = os.path.join(r"D:\Alpha", "Scripts")
if os.path.exists(fallback_scripts) and fallback_scripts not in sys.path:
    sys.path.insert(0, fallback_scripts)

from alpha_checks import get_active_tubex_file, check_not_locked, atomic_save
from add_order import get_mrp_boundaries, rearrange_mrp_orders

MONTH_MAP = {
    'jan': 'Jan', 'january': 'Jan', '1': 'Jan', '01': 'Jan',
    'feb': 'Feb', 'february': 'Feb', '2': 'Feb', '02': 'Feb',
    'mar': 'Mar', 'march': 'Mar', '3': 'Mar', '03': 'Mar',
    'apr': 'Apr', 'april': 'Apr', '4': 'Apr', '04': 'Apr',
    'may': 'May', '5': 'May', '05': 'May',
    'jun': 'Jun', 'june': 'Jun', '6': 'Jun', '06': 'Jun',
    'jul': 'Jul', 'july': 'Jul', '7': 'Jul', '07': 'Jul',
    'aug': 'Aug', 'august': 'Aug', '8': 'Aug', '08': 'Aug',
    'sep': 'Sep', 'september': 'Sep', '9': 'Sep', '09': 'Sep',
    'oct': 'Oct', 'october': 'Oct', '10': 'Oct',
    'nov': 'Nov', 'november': 'Nov', '11': 'Nov',
    'dec': 'Dec', 'december': 'Dec', '12': 'Dec',
}


def parse_month_year(month_input, year_input):
    mon_key = str(month_input).strip().lower()
    mon_title = MONTH_MAP.get(mon_key) or (MONTH_MAP.get(mon_key[:3]) if len(mon_key) >= 3 else None)
    if not mon_title:
        raise ValueError(f"Unrecognized month: '{month_input}'")

    y = int(year_input)
    if y < 100:
        y = 2000 + y
    yy_str = f"{y % 100:02d}"

    return mon_title, y, yy_str


def safe_eval_math(expr):
    if expr is None:
        return 0
    if isinstance(expr, (int, float)):
        return int(round(expr))
    expr_str = str(expr).lstrip('=').strip()
    if re.match(r'^[0-9+\-*/().\s]+$', expr_str):
        try:
            return int(round(float(eval(expr_str))))
        except Exception:
            return 0
    return 0


def compute_mtd_production(ws_pl):
    mtd = {}
    for row in ws_pl.iter_rows(min_row=3, values_only=True):
        if len(row) < 8:
            continue
        machine   = row[1]   # col B
        prod_name = row[3]   # col D
        pid       = row[5]   # col F
        good_qty  = row[7]   # col H

        if not machine or not pid or not good_qty:
            continue

        mach_up  = str(machine).upper()
        is_print = mach_up.startswith('PRINT') or mach_up.startswith('PLINE')
        is_pet   = mach_up.startswith('PF') or mach_up.startswith('PET')
        is_varn  = '(VARNISH)' in str(prod_name).upper()

        if (is_print and not is_varn) or is_pet:
            try:
                pid_int = int(pid)
                mtd[pid_int] = mtd.get(pid_int, 0) + int(good_qty)
            except (TypeError, ValueError):
                pass
    return mtd


def roll_new_month(month_str, year_val, dry_run=False, root_dir=None):
    if root_dir is None:
        cwd = os.getcwd()
        if get_active_tubex_file(cwd):
            root_dir = os.path.abspath(cwd)
        else:
            root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

    mon_title, y, yy_str = parse_month_year(month_str, year_val)
    new_filename = f"Tubex_{mon_title}{yy_str}.xlsx"
    new_filepath = os.path.join(root_dir, new_filename)

    # Step 1: Confirm no Tubex_<NewMon><NewYY>.xlsx already exists in root
    if os.path.exists(new_filepath):
        print(f"[ERROR] Target workbook '{new_filename}' already exists in root: {new_filepath}")
        print("        Aborting rollover to prevent overwriting.")
        sys.exit(1)

    # Find active Tubex workbook
    active_file = get_active_tubex_file(root_dir)
    if not active_file or not os.path.exists(active_file):
        print(f"[ERROR] No active Tubex*.xlsx workbook found in {root_dir}.")
        sys.exit(1)

    check_not_locked(active_file)

    sandbox_dir = None
    if dry_run:
        print("\n" + "=" * 70)
        print("   [DRY-RUN MODE ACTIVE] Operating on a temporary sandbox copy.")
        print("   Live production workbook will NOT be modified.")
        print("=" * 70 + "\n")
        sandbox_dir = tempfile.mkdtemp(prefix="dry_run_new_month_")
        sandbox_active = os.path.join(sandbox_dir, os.path.basename(active_file))
        shutil.copy2(active_file, sandbox_active)
        sandbox_records = os.path.join(sandbox_dir, "Tubex Records")
        os.makedirs(sandbox_records, exist_ok=True)
        working_root = sandbox_dir
        working_active = sandbox_active
        working_new = os.path.join(sandbox_dir, new_filename)
    else:
        working_root = root_dir
        working_active = active_file
        working_new = new_filepath

    # Step 2: Copy current active workbook to new filename in root
    shutil.copy2(working_active, working_new)

    # Step 3: In the new copy's MRP sheet: compute remaining = Required Qty - Produced
    # Load formula workbook to read formulas and evaluate quantities reliably
    wb_src = openpyxl.load_workbook(working_active, data_only=False)
    ws_src_mrp = wb_src['MRP']
    b_val = get_mrp_boundaries(ws_src_mrp)

    # Compute MTD production from Production_Log
    mtd_prod = {}
    if 'Production_Log' in wb_src.sheetnames:
        mtd_prod = compute_mtd_production(wb_src['Production_Log'])

    # Also load cached values if available
    val_map = {}
    try:
        wb_val = openpyxl.load_workbook(working_active, data_only=True, read_only=True)
        ws_val_mrp = wb_val['MRP']
        for sec, s, e in [('TUBE', b_val['tube_start'], b_val['tube_end']), ('PET', b_val['pet_start'], b_val['pet_end'])]:
            for r in range(s, e + 1):
                pid = ws_val_mrp.cell(r, 4).value
                if pid is not None:
                    try:
                        val_map[int(pid)] = {
                            'req': ws_val_mrp.cell(r, 6).value,
                            'prod': ws_val_mrp.cell(r, 7).value,
                        }
                    except (ValueError, TypeError):
                        pass
        wb_val.close()
    except Exception:
        pass

    order_decisions = []
    discrepancies = []
    for sec, s, e in [('TUBE', b_val['tube_start'], b_val['tube_end']), ('PET', b_val['pet_start'], b_val['pet_end'])]:
        for r in range(s, e + 1):
            pid = ws_src_mrp.cell(r, 4).value
            if pid is None:
                continue
            try:
                pid = int(pid)
            except (ValueError, TypeError):
                continue

            name = ws_src_mrp.cell(r, 3).value
            jof = ws_src_mrp.cell(r, 5).value

            raw_req = ws_src_mrp.cell(r, 6).value
            req = safe_eval_math(raw_req)
            if req == 0 and pid in val_map and val_map[pid]['req'] is not None:
                req = safe_eval_math(val_map[pid]['req'])

            prod = None
            if pid in val_map and val_map[pid]['prod'] is not None:
                v = val_map[pid]['prod']
                if not str(v).startswith('='):
                    try:
                        prod = int(round(v))
                    except (ValueError, TypeError):
                        pass

            prod_from_log = mtd_prod.get(pid, 0)
            if prod is None:
                prod = prod_from_log
            else:
                if abs(prod - prod_from_log) > 1:
                    discrepancies.append({
                        'pid': pid,
                        'name': name,
                        'cached_prod': prod,
                        'prod_from_log': prod_from_log,
                        'chosen': max(prod, prod_from_log),
                    })
                prod = max(prod, prod_from_log)

            rem = req - prod

            order_decisions.append({
                'sec': sec,
                'row': r,
                'pid': pid,
                'name': name,
                'jof': jof,
                'req': req,
                'prod': prod,
                'rem': rem,
            })
    wb_src.close()

    # Apply order decisions and clear logs in new workbook
    wb = openpyxl.load_workbook(working_new, data_only=False)
    ws_mrp = wb['MRP']

    carried_forward = []
    finished_orders = []
    rows_to_delete = []

    for dec in order_decisions:
        r = dec['row']
        rem = dec['rem']
        if rem > 0:
            ws_mrp.cell(r, 6).value = rem
            carried_forward.append(dec)
        else:
            rows_to_delete.append(r)
            finished_orders.append(dec)

    # Delete finished order rows from bottom to top
    for r in sorted(rows_to_delete, reverse=True):
        ws_mrp.delete_rows(r)

    # Reflow formulas using add_order.py logic
    rearrange_mrp_orders(ws_mrp)

    # Step 4: Clear Production_Log to header rows (rows 1-2)
    # Clear FG Stock to header rows (rows 1-3)
    if 'Production_Log' in wb.sheetnames:
        ws_pl = wb['Production_Log']
        if ws_pl.max_row >= 3:
            ws_pl.delete_rows(3, ws_pl.max_row - 2)

    if 'FG Stock' in wb.sheetnames:
        ws_fg = wb['FG Stock']
        if ws_fg.max_row >= 4:
            ws_fg.delete_rows(4, ws_fg.max_row - 3)

    # Step 5 & 6 prep: Reset Dashboard orders/dispatch references
    if 'Tubex_Dashboard' in wb.sheetnames:
        ws_dash = wb['Tubex_Dashboard']
        for r in range(11, ws_dash.max_row + 1):
            pid_val = ws_dash.cell(r, 6).value
            if pid_val:
                ws_dash.cell(r, 7).value = f'=IFERROR(INDEX(MRP!$F$3:$F$150,MATCH(F{r},MRP!$D$3:$D$150,0)),0)'
                ws_dash.cell(r, 11).value = None

    try:
        atomic_save(wb, working_new)
    except Exception:
        wb.save(working_new)
    wb.close()

    # Step 6: Run sort_dashboard.py against the new workbook
    sort_script = os.path.join(SCRIPTS_DIR, "sort_dashboard.py")
    if not os.path.exists(sort_script):
        sort_script = os.path.join(fallback_scripts, "sort_dashboard.py")

    if os.path.exists(sort_script):
        env = os.environ.copy()
        env["TUBEX_FILE"] = os.path.abspath(working_new)
        rc = subprocess.run([sys.executable, sort_script], env=env, capture_output=True, text=True)
        if rc.returncode != 0:
            print(f"[ERROR] sort_dashboard.py failed on {working_new}:\n{rc.stderr or rc.stdout}")
            sys.exit(1)

    # Step 7: Move OLD month's workbook into "Tubex Records\"
    records_dir = os.path.join(working_root, "Tubex Records")
    os.makedirs(records_dir, exist_ok=True)
    old_base = os.path.basename(active_file if not dry_run else working_active)
    archived_path = os.path.join(records_dir, old_base)

    if not dry_run:
        shutil.move(active_file, archived_path)
        old_location_str = archived_path
    else:
        shutil.move(working_active, archived_path)
        old_location_str = os.path.join(root_dir, "Tubex Records", old_base)

    # Step 8: Print summary
    print("\n" + "=" * 75)
    print(f"   NEW MONTH ROLLOVER SUMMARY: {mon_title} {y}")
    print("=" * 75)
    if discrepancies:
        print(f"\n[WARNING] Production Count Discrepancies ({len(discrepancies)}):")
        for d in discrepancies:
            print(f"  ! PID {d['pid']} ({d['name']}): cached MRP Produced = {d['cached_prod']:,}, Production_Log = {d['prod_from_log']:,} - used the higher of the two ({d['chosen']:,}), verify this one manually")

    print(f"\nCarried-Forward Orders ({len(carried_forward)}):")
    if carried_forward:
        for item in carried_forward:
            jof_str = str(item['jof']) if item['jof'] else '-'
            print(f"  • [{item['sec']:<4}] PID {item['pid']:<5} | {item['name'][:32]:<32} | New Req: {item['rem']:>8,d} (was {item['req']:,}, prod {item['prod']:,}) | JOF: {jof_str}")
    else:
        print("  (None)")

    print(f"\nFinished Orders Dropped ({len(finished_orders)}):")
    if finished_orders:
        for item in finished_orders:
            jof_str = str(item['jof']) if item['jof'] else '-'
            print(f"  • [{item['sec']:<4}] PID {item['pid']:<5} | {item['name'][:32]:<32} | Req: {item['req']:>8,d} | Prod: {item['prod']:>8,d} | Rem: {item['rem']:>8,d} | JOF: {jof_str}")
    else:
        print("  (None)")

    print("\nWorkbook Locations:")
    if not dry_run:
        print(f"  • New Active Workbook:   {new_filepath}")
        print(f"  • Old Workbook Archived: {old_location_str}")
    else:
        print(f"  • [DRY-RUN] New Workbook would be:   {new_filepath}")
        print(f"  • [DRY-RUN] Old Workbook would move: {old_location_str}")
        try:
            shutil.rmtree(sandbox_dir)
        except Exception:
            pass

    print("=" * 75)
    return True


def main():
    parser = argparse.ArgumentParser(description="Roll Tubex workbook over to a new month.")
    parser.add_argument('--month', '-m', required=True, help="New month (e.g. Oct, October, 10)")
    parser.add_argument('--year', '-y', required=True, type=int, help="New year (e.g. 2026, 26)")
    parser.add_argument('--dry-run', action='store_true', help="Run in a temporary sandbox without modifying live files")

    args = parser.parse_args()
    roll_new_month(args.month, args.year, dry_run=args.dry_run)


if __name__ == '__main__':
    main()
